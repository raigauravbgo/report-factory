import io
import logging
import time
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

logger = logging.getLogger(__name__)

from core.database import get_db
from models.report_recipe import ReportRecipe
from models.staging_table import StagingTable
from models.upload import Upload
from services.compute import compute_dashboard, get_filter_options

router = APIRouter(prefix="/api/dashboard", tags=["dashboard"])


@router.get("")
def list_dashboards(db: Session = Depends(get_db)):
    """List all generated dashboards (ReportRecipes) for the Library page.

    Returns each recipe with the file names and KPI names from its config so
    the Library page can display them without a separate data fetch.
    """
    from models.dataset import Dataset

    recipes = (
        db.query(ReportRecipe)
        .order_by(ReportRecipe.created_at.desc())
        .all()
    )
    result = []
    for r in recipes:
        cfg = r.config or {}
        dataset_id = r.dataset_id or cfg.get("dataset_id")

        # Collect file names for this dataset (skip virtual dimension)
        filenames: list[str] = []
        if dataset_id:
            uploads = (
                db.query(Upload)
                .filter(
                    Upload.dataset_id == dataset_id,
                    Upload.filename != "__virtual_dimension__",
                )
                .all()
            )
            filenames = [u.filename for u in uploads]

        kpi_names = [k.get("name", "") for k in cfg.get("kpis", [])]

        result.append({
            "recipe_id": r.id,
            "dataset_id": dataset_id,
            "client_id": r.client_id or "—",
            "filenames": filenames,
            "kpi_names": kpi_names,
            "kpi_count": len(kpi_names),
            "date_column": cfg.get("date_column") or "",
            "granularity": cfg.get("granularity") or "weekly",
            "created_at": r.created_at.isoformat() if r.created_at else None,
        })
    return result


def _get_recipe_and_staging(recipe_id: int, db: Session):
    """Return (recipe, config, primary_staging, all_table_names).

    all_table_names includes every staging table for the dataset so that
    multi-file uploads (e.g. Adherence + CSAT) are concatenated in compute.
    """
    recipe = db.query(ReportRecipe).filter(ReportRecipe.id == recipe_id).first()
    if not recipe:
        raise HTTPException(404, "Recipe not found.")
    config = recipe.config
    upload_id = config.get("upload_id")
    dataset_id = config.get("dataset_id")

    # Primary staging table (backward compat — must resolve)
    primary: StagingTable | None = None
    if upload_id:
        primary = db.query(StagingTable).filter(StagingTable.upload_id == upload_id).first()

    # Collect ALL staging tables for the dataset to support multi-file compute.
    # If dataset_id is present use it; otherwise fall back to single upload.
    # H4: Single JOIN query replaces N+1 (one StagingTable query per upload)
    # Exclude virtual dimension tables — they are built from fact-table data and
    # are only used for filter enrichment. Including them in the concat doubles
    # row counts for every count/sum KPI (e.g. 17 503 rows → 35 006).
    all_table_names: list[str] = []
    if dataset_id:
        staging_rows = (
            db.query(StagingTable)
            .join(Upload, StagingTable.upload_id == Upload.id)
            .filter(Upload.dataset_id == dataset_id)
            .filter(Upload.filename != "__virtual_dimension__")
            .all()
        )
        all_table_names = [st.table_name for st in staging_rows]

    if not all_table_names and primary:
        all_table_names = [primary.table_name]

    if not primary and all_table_names:
        primary = db.query(StagingTable).filter(
            StagingTable.table_name == all_table_names[0]
        ).first()

    if not primary:
        raise HTTPException(404, "Staging data not found. Re-upload the file.")

    return recipe, config, primary, all_table_names


@router.get("/{recipe_id}/validate-config")
def validate_config(recipe_id: int, db: Session = Depends(get_db)):
    """
    Pre-flight data integrity check: verifies the recipe config works against real data.
    Returns { valid, errors, warnings, details }.
    Used by the recipe editor before approval and by the dashboard on first load.
    """
    recipe, config, _staging, all_table_names = _get_recipe_and_staging(recipe_id, db)
    from services.compute import validate_dashboard_config
    result = validate_dashboard_config(config, all_table_names)
    logger.info(
        "VALIDATE_CONFIG recipe_id=%d valid=%s errors=%d warnings=%d",
        recipe_id, result["valid"], len(result["errors"]), len(result["warnings"]),
    )
    return result


@router.get("/{recipe_id}/filter-values")
def get_dashboard_filter_values(recipe_id: int, request: Request, db: Session = Depends(get_db)):
    """Return distinct values for each filter column — used to populate FilterBar dropdowns.

    Supports cascading (Power BI-style) filters: pass currently-active filter values as query
    params (same format as the /data endpoint).  Each column's options are computed from the
    dataset with ALL OTHER active filters applied, so selecting location=India narrows the
    department dropdown to departments that exist within India.
    """
    recipe, config, staging, all_table_names = _get_recipe_and_staging(recipe_id, db)

    # Extract active filters from query params (same convention as /data endpoint).
    reserved = {"recipe_id", "granularity"}
    active_filters = {
        k: v for k, v in request.query_params.items()
        if k not in reserved and v
    }

    # Only show columns explicitly listed as filters; fall back to dimensions if none set.
    filter_cols = config.get("filters") or []
    candidate_cols = filter_cols if filter_cols else config.get("dimensions", [])
    candidate_cols = list(dict.fromkeys(candidate_cols))  # deduplicate, preserve order

    # Exclude entity keys, emails, and text blobs (high-cardinality / PII) from dropdowns.
    # Bug fix: build col_meta from ALL files in the dataset, not just the primary staging table.
    # Columns from files 2+ would otherwise bypass semantic-tag filtering.
    dataset_id = config.get("dataset_id")
    all_col_meta: dict = {}
    if dataset_id:
        all_stagings = (
            db.query(StagingTable)
            .join(Upload, StagingTable.upload_id == Upload.id)
            .filter(Upload.dataset_id == dataset_id)
            .all()
        )
        for _st in all_stagings:
            for _col in (_st.profile_data or {}).get("columns", []):
                all_col_meta[_col["name"]] = _col  # last file wins on name collision
    elif staging.profile_data:
        # fallback for legacy single-file recipes without dataset_id
        for _col in staging.profile_data.get("columns", []):
            all_col_meta[_col["name"]] = _col
    if all_col_meta:
        filtered = [
            c for c in candidate_cols
            if all_col_meta.get(c, {}).get("semantic_tag") not in ("entity_key", "time_key", "text", "ignore")
        ]
        candidate_cols = filtered if filtered else candidate_cols

    # Cap dimension-fallback columns at 5 by unique_count ascending (lowest cardinality = most
    # useful for filtering). Only applies when no explicit filter columns are configured —
    # user-chosen filters are never capped.
    if not filter_cols and len(candidate_cols) > 5:
        def _unique_count(col_name: str) -> int:
            return int(all_col_meta.get(col_name, {}).get("unique_count", 999))
        candidate_cols = sorted(candidate_cols, key=_unique_count)[:5]
        logger.debug("FILTER_BAR_CAP recipe_id=%d capped_dims=%s", recipe_id, candidate_cols)

    try:
        options = get_filter_options(
            all_table_names,
            candidate_cols,
            confirmed_relationships=config.get("confirmed_relationships") or None,
            upload_table_map=config.get("upload_table_map") or None,
            active_filters=active_filters or None,
            date_column=config.get("date_column") or None,
        )
    except Exception as e:
        raise HTTPException(500, f"Failed to load filter values: {e}")
    return options


@router.get("/{recipe_id}/data")
def get_dashboard_data(
    recipe_id: int,
    request: Request,
    db: Session = Depends(get_db),
    granularity: str = Query(default=None, description="Override recipe granularity: daily | weekly | monthly"),
):
    recipe, config, staging, all_table_names = _get_recipe_and_staging(recipe_id, db)

    # Extract active filters from query params (exclude reserved params)
    reserved = {"recipe_id", "granularity"}
    active_filters = {
        k: v for k, v in request.query_params.items()
        if k not in reserved and v
    }

    # For multi-file recipes generated before the sessionStorage fix, confirmed_relationships
    # may be empty even though detectable JOINs exist.  Auto-infer them at query time so
    # row-level filters produce correct results instead of nulling out all KPI values.
    if (
        active_filters
        and not config.get("confirmed_relationships")
        and len(all_table_names) > 1
        and config.get("dataset_id")
        and config.get("upload_table_map")
    ):
        from services.schema_relationships import infer as _infer_rels
        try:
            inferred = _infer_rels(config["dataset_id"], db)
            pk_fk = [
                {
                    "relationship_type": r.relationship_type,
                    "file_a": r.file_a,
                    "col_a": r.col_a,
                    "file_b": r.file_b,
                    "col_b": r.col_b,
                    "confidence": r.confidence,
                }
                for r in inferred
                if r.relationship_type == "pk_fk" and r.confidence >= 0.7
            ]
            if pk_fk:
                config = {**config, "confirmed_relationships": pk_fk}
                logger.info(
                    "AUTO_REL recipe_id=%d inferred %d pk_fk relationships for filtered query",
                    recipe_id,
                    len(pk_fk),
                )
        except Exception as _e:
            logger.warning("AUTO_REL recipe_id=%d failed to infer: %s", recipe_id, _e)

    t0 = time.perf_counter()
    try:
        result = compute_dashboard(
            config, all_table_names,
            filters=active_filters or None,
            granularity_override=granularity,
        )
    except Exception as e:
        logger.error("COMPUTE_FAIL recipe_id=%d error=%r", recipe_id, str(e))
        raise HTTPException(500, f"Computation failed: {e}")
    logger.info(
        "DASHBOARD_COMPUTE recipe_id=%d kpis=%d blank=%d filters=%s duration=%.2fs",
        recipe_id,
        len(result.get("kpi_summaries", [])),
        sum(1 for k in result.get("kpi_summaries", []) if k.get("value") is None),
        list(active_filters.keys()) if active_filters else [],
        time.perf_counter() - t0,
    )
    effective_granularity = granularity or config.get("granularity", "weekly")

    # ── Enhanced View enrichment fields (additive — Classic fields unchanged) ──

    # time_series_dict: Record<kpiName, [{period, value}]>
    ts_dict: dict = {
        ts["kpi"]: [{"period": p["date"], "value": p["value"]} for p in ts["data"]]
        for ts in result.get("time_series", [])
    }

    # dimension_breakdowns: reshape breakdown array → {dim: {kpiName: [{name, value}]}}
    dim_breakdowns: dict = {}
    for bk in result.get("breakdown", []):
        dim = bk["dimension"]
        kpi_name = bk["kpi"]
        if dim not in dim_breakdowns:
            dim_breakdowns[dim] = {}
        dim_breakdowns[dim][kpi_name] = [
            {"name": str(p["label"]), "value": float(p["value"])}
            for p in bk["data"]
        ]

    # dimension_spreads: (max - min) / mean per dimension, averaged across KPIs
    def _spread_for_dim(kpis_data: dict) -> float:
        kpi_spreads = []
        for segments in kpis_data.values():
            values = [s["value"] for s in segments]
            if len(values) < 2:
                continue
            mean_val = sum(values) / len(values)
            if mean_val == 0:
                continue
            kpi_spreads.append((max(values) - min(values)) / mean_val)
        return sum(kpi_spreads) / len(kpi_spreads) if kpi_spreads else 0.0

    dim_spreads: dict = {
        dim: round(_spread_for_dim(kpis_data), 4)
        for dim, kpis_data in dim_breakdowns.items()
    }
    dim_breakdowns = dict(
        sorted(dim_breakdowns.items(), key=lambda x: dim_spreads.get(x[0], 0), reverse=True)
    )

    # metrics: richer KPI objects with delta, status, direction, prior_value
    _lower_keywords = {"wait", "handle_time", "aht", "error", "abandon", "escalat",
                       "delay", "cost", "churn", "attrition", "shrinkage", "overtime"}
    ts_by_kpi: dict = {ts["kpi"]: ts["data"] for ts in result.get("time_series", [])}
    metrics: list = []
    for i, kpi in enumerate(result.get("kpi_summaries", [])):
        kpi_ts = ts_by_kpi.get(kpi["name"], [])
        valid_ts = [p for p in kpi_ts if p.get("value") is not None]
        current_val = kpi["value"]
        prior_val = valid_ts[-2]["value"] if len(valid_ts) >= 2 else None
        period_label = valid_ts[-1]["date"] if valid_ts else None

        if current_val is not None and prior_val is not None and prior_val != 0:
            delta = current_val - prior_val
            delta_pct = (delta / abs(prior_val)) * 100
        else:
            delta = None
            delta_pct = None

        name_lower = kpi["name"].lower()
        direction = (
            "lower_is_better"
            if any(kw in name_lower for kw in _lower_keywords)
            else "higher_is_better"
        )

        if delta is None or delta_pct is None:
            status = "neutral"
        else:
            up = delta >= 0
            good = (not up) if direction == "lower_is_better" else up
            abs_pct = abs(delta_pct)
            if abs_pct < 5:
                status = "neutral"
            elif good and abs_pct >= 5:
                status = "good"
            elif not good and abs_pct >= 20:
                status = "risk"
            else:
                status = "warning"

        metrics.append({
            "id": i,
            "name": kpi["name"],
            "value": round(float(current_val), 4) if current_val is not None else 0.0,
            "delta": round(float(delta), 4) if delta is not None else None,
            "delta_pct": round(float(delta_pct), 2) if delta_pct is not None else None,
            "status": status,
            "direction": direction,
            "prior_value": round(float(prior_val), 4) if prior_val is not None else None,
            "count": len(valid_ts),
            "period": period_label,
            "formula": kpi["formula"],
            "format": kpi.get("format", ""),
        })

    # data_quality: derive from time_series coverage
    all_dates = sorted({
        p["date"]
        for ts in result.get("time_series", [])
        for p in ts["data"]
        if p.get("date")
    })
    dq_warnings: list[str] = []
    blank_kpis = [k["name"] for k in result.get("kpi_summaries", []) if k.get("value") is None]
    if blank_kpis:
        dq_warnings.append(f"No data for: {', '.join(blank_kpis)}")
    data_quality = {
        "status": "warning" if dq_warnings else "ok",
        "date_coverage": f"{all_dates[0]} to {all_dates[-1]}" if len(all_dates) >= 2 else (all_dates[0] if all_dates else None),
        "most_recent_date": all_dates[-1] if all_dates else None,
        "warnings": dq_warnings,
    }

    # row_count: quick COUNT(*) from staging tables
    from sqlalchemy import text as _sql_text
    row_count = 0
    try:
        for tname in all_table_names:
            row_count += db.execute(_sql_text(f'SELECT COUNT(*) FROM "{tname}"')).scalar() or 0
    except Exception:
        row_count = sum(len(ts["data"]) for ts in result.get("time_series", []))

    # approved flag
    approved = recipe.approved_at is not None

    # insights enriched with driver and impact (synthetic defaults when absent)
    _sev_impact = {
        "critical": "Significant performance risk requiring immediate attention",
        "high": "High impact on KPI performance — review urgently",
        "medium": "Moderate impact on operational metrics",
        "low": "Minor variance within acceptable range",
    }
    enriched_insights: list = []
    for ins in result.get("insights", []):
        enriched = dict(ins)
        if "driver" not in enriched or not enriched.get("driver"):
            enriched["driver"] = (ins.get("finding") or "")[:120] or ins.get("headline", "")
        if "impact" not in enriched or not enriched.get("impact"):
            enriched["impact"] = _sev_impact.get(ins.get("severity", "low"), _sev_impact["low"])
        enriched_insights.append(enriched)

    # Synthesise a per-KPI insight for any metric the AI didn't cover
    _status_to_sev = {"risk": "high", "warning": "medium", "good": "low", "neutral": "low"}
    _covered_headlines = {ins.get("headline", "").lower() for ins in enriched_insights}
    for m in metrics:
        mname = m["name"]
        if any(mname.lower() in h for h in _covered_headlines):
            continue
        val = m["value"]
        sev = _status_to_sev.get(m.get("status", "neutral"), "low")
        delta_pct = m.get("delta_pct")
        prior = m.get("prior_value")
        if delta_pct is not None and prior is not None:
            went_up = delta_pct >= 0
            hib = m.get("direction") == "higher_is_better"
            direction_word = "improved" if (went_up == hib) else "declined"
            if abs(delta_pct) < 0.1:
                finding = f"No significant change from prior period. Current: {val:.1f}, Prior: {prior:.1f}."
            else:
                finding = f"{direction_word.capitalize()} {abs(delta_pct):.1f}% from prior period. Current: {val:.1f}, Prior: {prior:.1f}."
        else:
            finding = f"Current value: {val:.1f}."
        enriched_insights.append({
            "severity": sev,
            "headline": f"{mname} at {val:.1f}",
            "finding": finding,
            "driver": finding[:120],
            "impact": _sev_impact.get(sev, _sev_impact["low"]),
            "action": None,
        })

    # ── end enrichment ────────────────────────────────────────────────────────

    return {
        "recipe_id": recipe_id,
        "config": {**config, "granularity": effective_granularity},
        "active_filters": active_filters,
        "active_granularity": effective_granularity,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        # Classic fields (unchanged)
        **result,
        # Enhanced View fields (additive)
        "metrics": metrics,
        "time_series_dict": ts_dict,
        "dimension_breakdowns": dim_breakdowns,
        "dimension_spreads": dim_spreads,
        "data_quality": data_quality,
        "row_count": row_count,
        "approved": approved,
        "insights": enriched_insights,
    }


@router.get("/{recipe_id}/export/excel")
def export_excel(recipe_id: int, db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    recipe, config, staging, all_table_names = _get_recipe_and_staging(recipe_id, db)
    try:
        result = compute_dashboard(config, all_table_names)
    except Exception as e:
        logger.error("EXPORT_FAIL recipe_id=%d format=excel error=%r", recipe_id, str(e))
        raise HTTPException(500, f"Computation failed: {e}")
    logger.info("DASHBOARD_EXPORT recipe_id=%d format=excel kpis=%d", recipe_id, len(result.get("kpi_summaries", [])))

    wb = openpyxl.Workbook()

    # Sheet 1: KPI Summary
    ws = wb.active
    ws.title = "KPI Summary"
    header_fill = PatternFill("solid", fgColor="3B82F6")
    header_font = Font(bold=True, color="FFFFFF")

    ws.append(["KPI", "Value", "Formula"])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for kpi in result["kpi_summaries"]:
        val = kpi["value"]
        if val is not None and "/" in kpi["formula"]:
            val = f"{val * 100:.1f}%"
        elif val is not None:
            val = round(val, 2)
        ws.append([kpi["name"], val, kpi["formula"]])

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 35

    # Sheet 2: Time Series
    if result["time_series"]:
        ws2 = wb.create_sheet("Time Series")
        for ts in result["time_series"]:
            ws2.append([ts["kpi"]])
            ws2.cell(ws2.max_row, 1).font = Font(bold=True)
            ws2.append(["Date", "Value"])
            for cell in ws2[ws2.max_row]:
                cell.fill = header_fill
                cell.font = header_font
            for point in ts["data"]:
                ws2.append([point["date"], point["value"]])
            ws2.append([])
        ws2.column_dimensions["A"].width = 18
        ws2.column_dimensions["B"].width = 15

    # Sheet 3: Breakdown
    if result["breakdown"]:
        ws3 = wb.create_sheet("Breakdown")
        for bk in result["breakdown"]:
            ws3.append([f"{bk['kpi']} by {bk['dimension']}"])
            ws3.cell(ws3.max_row, 1).font = Font(bold=True)
            ws3.append([bk["dimension"], "Value"])
            for cell in ws3[ws3.max_row]:
                cell.fill = header_fill
                cell.font = header_font
            for point in bk["data"]:
                ws3.append([point["label"], point["value"]])
            ws3.append([])
        ws3.column_dimensions["A"].width = 25
        ws3.column_dimensions["B"].width = 15

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=dashboard_recipe_{recipe_id}.xlsx"},
    )


@router.get("/{recipe_id}/export/pptx")
def export_pptx(recipe_id: int, db: Session = Depends(get_db)):
    from exporters.pptx_exporter import generate_pptx

    recipe, config, staging, all_table_names = _get_recipe_and_staging(recipe_id, db)
    try:
        result = compute_dashboard(config, all_table_names)
    except Exception as e:
        logger.error("EXPORT_FAIL recipe_id=%d format=pptx error=%r", recipe_id, str(e))
        raise HTTPException(500, f"Computation failed: {e}")
    logger.info("DASHBOARD_EXPORT recipe_id=%d format=pptx kpis=%d", recipe_id, len(result.get("kpi_summaries", [])))

    dashboard_data = {
        "template_type": config.get("granularity", ""),
        "client_id": f"Recipe #{recipe_id}",
        "period_start": None,
        "period_end": None,
        "kpi_tiles": [
            {
                "kpi_id": k["name"],
                "display_name": k["name"].replace("_", " ").title(),
                "value": k["value"],
                "format": "percentage" if "/" in k["formula"] else "integer",
                "flags": [],
            }
            for k in result["kpi_summaries"]
        ],
        "charts": [
            {**ts, "type": "line"}
            for ts in result["time_series"]
        ] + [
            {**bk, "type": "bar", "title": f"{bk['kpi']} by {bk['dimension']}"}
            for bk in result["breakdown"]
        ],
        "data_quality_flags": [],
    }

    pptx_bytes = generate_pptx(dashboard_data)
    return StreamingResponse(
        io.BytesIO(pptx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f"attachment; filename=dashboard_recipe_{recipe_id}.pptx"},
    )


# ── Formula validation ───────────────────────────────────────────────────────

@router.post("/{recipe_id}/validate-formula")
def validate_formula(recipe_id: int, body: dict, db: Session = Depends(get_db)):
    """
    Test-compute a single KPI formula against the recipe's staging data.

    Returns:
      { valid: bool, preview_value: float|null, error: str|null }

    Used by the recipe editor to give per-formula feedback before approval.
    """
    formula = (body.get("formula") or "").strip()
    if not formula:
        raise HTTPException(400, "formula is required")

    recipe, config, _staging, all_table_names = _get_recipe_and_staging(recipe_id, db)

    try:
        from services.compute import (
            _load_staging_df,
            _eval_formula,
            _formula_agg,
        )
        import pandas as pd

        df = _load_staging_df(all_table_names)
        if df.empty:
            return {"valid": False, "preview_value": None,
                    "error": "Staging table is empty — re-upload the file."}

        series = _eval_formula(df, formula)
        if series.empty or series.isna().all():
            return {"valid": False, "preview_value": None,
                    "error": "Formula references unknown or non-numeric columns."}

        agg = _formula_agg(formula)
        raw = series.mean() if agg in ("mean", "ratio") else series.sum()

        if pd.isna(raw):
            return {"valid": False, "preview_value": None,
                    "error": "Formula produces no valid values (all NaN)."}

        return {"valid": True, "preview_value": round(float(raw), 4), "error": None}
    except Exception as exc:
        logger.warning("VALIDATE_FORMULA recipe_id=%d formula=%r error=%r",
                       recipe_id, formula, str(exc))
        return {"valid": False, "preview_value": None, "error": str(exc)}


# ── Recipe config patch ───────────────────────────────────────────────────────

_PATCHABLE_FIELDS = {"granularity", "date_column", "dimensions", "filters", "kpis", "column_mappings"}


@router.patch("/{recipe_id}/config")
def patch_recipe_config(recipe_id: int, body: dict, db: Session = Depends(get_db)):
    """Update mutable recipe config fields without changing approval status.

    Only granularity, date_column, dimensions, filters, kpis, and column_mappings
    are accepted — structural fields (upload_id, dataset_id) are immutable.
    """
    from models.report_recipe import ReportRecipe

    recipe = db.query(ReportRecipe).filter(ReportRecipe.id == recipe_id).first()
    if not recipe:
        raise HTTPException(404, "Recipe not found.")

    config = dict(recipe.config)
    updated_fields = []
    for field in _PATCHABLE_FIELDS:
        if field in body:
            config[field] = body[field]
            updated_fields.append(field)

    if not updated_fields:
        raise HTTPException(400, f"No patchable fields found. Allowed: {sorted(_PATCHABLE_FIELDS)}")

    recipe.config = config
    db.commit()
    logger.info("RECIPE_CONFIG_PATCHED recipe_id=%d fields=%s", recipe_id, updated_fields)
    return {"status": "ok", "updated": updated_fields, "config": config}
