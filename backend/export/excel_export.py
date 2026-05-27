from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=16, color="1F4E79")


def generate(kpi_results: dict, output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard Summary"

    ws["A1"] = "Executive Dashboard — AI-Generated Report"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:E1")
    ws.row_dimensions[1].height = 28

    headers = ["KPI ID", "Name", "Value", "Unit", "Chart Type"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    row = 4
    for kpi_id, result in kpi_results.items():
        if "error" in result:
            continue
        ws.cell(row=row, column=1, value=kpi_id)
        ws.cell(row=row, column=2, value=result.get("name", ""))
        ws.cell(row=row, column=3, value=result.get("value") if result.get("value") is not None else "—")
        ws.cell(row=row, column=4, value=result.get("unit", ""))
        ws.cell(row=row, column=5, value=result.get("chart_type", ""))

        if result.get("breakdown") and isinstance(result["breakdown"], dict):
            breakdown_row = row + 2
            ws.cell(row=breakdown_row - 1, column=1, value=f"{result.get('name')} — Breakdown")
            ws.cell(row=breakdown_row - 1, column=1).font = Font(bold=True)
            for label, val in result["breakdown"].items():
                ws.cell(row=breakdown_row, column=1, value=str(label))
                ws.cell(row=breakdown_row, column=2, value=val)
                breakdown_row += 1
            row = breakdown_row + 1
        else:
            row += 1

    widths = [20, 30, 15, 10, 15]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(output_path)
