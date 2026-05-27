from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

_PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
_ACCENT = RGBColor(0x70, 0xAD, 0x47)


def generate(kpi_results: dict, output_path: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    _add_title_slide(prs)
    _add_kpi_summary_slide(prs, kpi_results)

    for kpi_id, result in kpi_results.items():
        if result.get("breakdown") and isinstance(result["breakdown"], dict):
            _add_breakdown_slide(prs, result)

    prs.save(output_path)


def _add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Executive Dashboard"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = _PRIMARY
    slide.placeholders[1].text = "AI-Generated Analysis Report"


def _add_kpi_summary_slide(prs: Presentation, kpi_results: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    heading = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.7))
    tf = heading.text_frame
    tf.text = "KPI Summary"
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = _PRIMARY

    x, y, w, h = 0.5, 1.1, 2.8, 1.6
    max_x = 12.5

    for kpi_id, result in kpi_results.items():
        if "error" in result:
            continue
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True

        name_para = tf.paragraphs[0]
        name_para.text = result.get("name", kpi_id)
        name_para.runs[0].font.size = Pt(10)
        name_para.runs[0].font.bold = True
        name_para.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        val_para = tf.add_paragraph()
        val = result.get("value")
        unit = result.get("unit", "")
        val_para.text = (str(val) if val is not None else "—") + (" " + unit if unit else "")
        val_para.runs[0].font.size = Pt(20)
        val_para.runs[0].font.bold = True
        val_para.runs[0].font.color.rgb = _PRIMARY

        x += w + 0.3
        if x + w > max_x:
            x = 0.5
            y += h + 0.3


def _add_breakdown_slide(prs: Presentation, result: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    heading = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.6))
    tf = heading.text_frame
    tf.text = result.get("name", result["kpi_id"])
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = _PRIMARY

    table_data = list(result["breakdown"].items())[:20]
    if not table_data:
        return

    rows = len(table_data) + 1
    table = slide.shapes.add_table(rows, 2, Inches(0.5), Inches(1.0), Inches(6), Inches(0.4 * rows)).table

    for header, text in zip(table.rows[0].cells, ["Label", "Value"]):
        header.text = text
        header.text_frame.paragraphs[0].runs[0].font.bold = True
        header.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, (label, val) in enumerate(table_data, start=1):
        table.rows[row_idx].cells[0].text = str(label)
        table.rows[row_idx].cells[1].text = str(val)
